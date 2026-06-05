"""
calculer_rating.py
──────────────────
Calcule le pollster rating de chaque institut à partir des sondages
terminaux compilés dans template_sondages.xlsx.

Trois facteurs de pondération sont combinés :
  1. Poids du scrutin    : défini manuellement dans B4 de chaque onglet
                           (ex : 1.0 pour présidentielle T1, 0.5 pour T2)
  2. Poids temporel      : décroissance exponentielle depuis ANNEE_REFERENCE
                           (scrutins anciens pèsent moins)
  3. Poids de proximité  : décroissance exponentielle selon le nombre de jours
                           entre la fin du terrain et le jour du scrutin
                           (sondages lointains pèsent moins)

Ce module est la SOURCE UNIQUE du calcul de rating : agregateur.py importe
la fonction `calculer_ratings()` au lieu de redupliquer la logique, ce qui
garantit des ratings identiques des deux côtés.

Usage en ligne de commande (affiche le détail) :
  python3 calculer_rating.py
"""

import math
import openpyxl
from collections import defaultdict
from datetime import datetime

# ══════════════════════════════════════════════════════════════
# PARAMÈTRES — modifier ici si besoin
# ══════════════════════════════════════════════════════════════

FICHIER          = 'template_sondages.xlsx'
ANNEE_REFERENCE  = datetime.today().year   # année courante : s'ajuste seul chaque année
DEMI_VIE_ANS     = 4      # demi-vie en années (poids temporel)
DEMI_VIE_JOURS   = 10     # demi-vie en jours  (poids de proximité)
NB_SCRUTINS_MIN  = 3      # scrutins minimum pour figurer dans le classement


# ══════════════════════════════════════════════════════════════
# FONCTIONS
# ══════════════════════════════════════════════════════════════

def normaliser_date(val):
    """Convertit n'importe quelle valeur de cellule date en objet datetime."""
    if isinstance(val, datetime):
        return val
    if hasattr(val, 'year'):          # objet date (sans heure)
        return datetime(val.year, val.month, val.day)
    if isinstance(val, str):
        try:
            parties = val.split('/')
            if len(parties) == 3:     # format DD/MM/YYYY
                j, m, a = parties
                return datetime(int(a), int(m), int(j))
            return datetime.fromisoformat(val[:10])
        except Exception:
            return None
    return None


def poids_temporel(annee_scrutin, annee_ref, demi_vie_ans):
    """0.5^(années_écoulées / demi_vie). Vaut 1.0 si annee_scrutin == annee_ref."""
    return 0.5 ** ((annee_ref - annee_scrutin) / demi_vie_ans)


def poids_proximite(date_terrain, date_scrutin, demi_vie_jours):
    """0.5^(jours_avant_scrutin / demi_vie). Vaut 1.0 si sondage le jour du scrutin."""
    if date_terrain is None or date_scrutin is None:
        return 1.0   # valeur neutre si date manquante
    jours = (date_scrutin - date_terrain).days
    jours = max(jours, 0)
    return 0.5 ** (jours / demi_vie_jours)


def lire_scrutin(ws, demi_vie_jours):
    """
    Lit un onglet et retourne un dict avec :
      - nom              : intitulé du scrutin
      - annee            : année du scrutin
      - date_scrutin     : objet datetime
      - poids            : poids du scrutin dans le rating (B4)
      - erreurs          : {nom_institut: (erreur_moy, poids_proximite)}
    Retourne None si l'onglet n'est pas lisible.
    """
    nom           = ws['B1'].value
    date_scrutin  = normaliser_date(ws['B2'].value)
    poids_scrutin = ws['B4'].value

    if date_scrutin is None or poids_scrutin is None:
        return None

    # Repérer les lignes structurelles
    res_row = ecart_row = None
    for r in range(1, ws.max_row + 1):
        val = ws.cell(r, 1).value
        if val == 'RÉSULTAT RÉEL':
            res_row = r
        if val == 'ÉCARTS ABSOLUS':
            ecart_row = r
            break

    if not res_row or not ecart_row:
        return None

    # Colonnes contenant les scores réels (un 0 signalerait un candidat non
    # mesuré, pas un score réel → on l'exclut comme dans l'agrégateur)
    cols_resultats = {
        c: ws.cell(res_row, c).value
        for c in range(4, ws.max_column + 1)
        if isinstance(ws.cell(res_row, c).value, (int, float))
        and ws.cell(res_row, c).value > 0
    }
    if not cols_resultats:
        return None

    # Erreur et proximité pour chaque institut
    erreurs = {}
    for r in range(res_row + 1, ecart_row):
        nom_inst = ws.cell(r, 1).value
        if not nom_inst or not isinstance(nom_inst, str):
            continue

        ecarts = [
            abs(ws.cell(r, c).value - val_reel)
            for c, val_reel in cols_resultats.items()
            if isinstance(ws.cell(r, c).value, (int, float))
            and ws.cell(r, c).value > 0
        ]
        if not ecarts:
            continue

        erreur_moy    = sum(ecarts) / len(ecarts)
        date_terrain  = normaliser_date(ws.cell(r, 2).value)
        prox          = poids_proximite(date_terrain, date_scrutin, demi_vie_jours)

        erreurs[nom_inst] = (erreur_moy, prox)

    return {
        'nom':         nom,
        'annee':       date_scrutin.year,
        'date_scrutin': date_scrutin,
        'poids':       poids_scrutin,
        'erreurs':     erreurs,
    }


def charger_scrutins(fichier, demi_vie_jours):
    """Lit tous les onglets d'un classeur et retourne la liste triée des scrutins."""
    wb = openpyxl.load_workbook(fichier, data_only=True)
    scrutins = []
    for sheet_name in wb.sheetnames:
        s = lire_scrutin(wb[sheet_name], demi_vie_jours)
        if s:
            scrutins.append(s)
    scrutins.sort(key=lambda s: s['annee'])
    return scrutins


def calculer_ratings(fichier=FICHIER, annee_ref=ANNEE_REFERENCE,
                     demi_vie_ans=DEMI_VIE_ANS, demi_vie_jours=DEMI_VIE_JOURS,
                     nb_scrutins_min=NB_SCRUTINS_MIN):
    """
    Calcule le rating de chaque institut.

    Retourne un quadruplet :
      ratings        : {institut: rating}  (rating = erreur_meilleur / erreur_inst)
      erreur_agregee : {institut: erreur moyenne pondérée}
      nb_scrutins    : {institut: nombre de scrutins couverts}
      scrutins       : liste des scrutins chargés (utile pour l'affichage)

    La normalisation se fait sur l'erreur du meilleur institut QUALIFIÉ
    (au moins `nb_scrutins_min` scrutins), pour qu'un institut vu une seule
    fois ne définisse pas l'échelle.
    """
    scrutins = charger_scrutins(fichier, demi_vie_jours)

    numerateur   = defaultdict(float)
    denominateur = defaultdict(float)
    nb_scrutins  = defaultdict(int)

    for s in scrutins:
        pt = poids_temporel(s['annee'], annee_ref, demi_vie_ans)
        for inst, (err, prox) in s['erreurs'].items():
            poids_total = s['poids'] * pt * prox
            numerateur[inst]   += err * poids_total
            denominateur[inst] += poids_total
            nb_scrutins[inst]  += 1

    erreur_agregee = {
        inst: numerateur[inst] / denominateur[inst]
        for inst in numerateur
        if denominateur[inst] > 0
    }

    instituts_qualifies = {
        inst: err
        for inst, err in erreur_agregee.items()
        if nb_scrutins[inst] >= nb_scrutins_min
    }
    if not instituts_qualifies:
        return {}, erreur_agregee, nb_scrutins, scrutins

    meilleure_erreur = min(instituts_qualifies.values())
    ratings = {
        inst: meilleure_erreur / err
        for inst, err in erreur_agregee.items()
    }
    return ratings, erreur_agregee, nb_scrutins, scrutins


# ══════════════════════════════════════════════════════════════
# AFFICHAGE EN LIGNE DE COMMANDE
# ══════════════════════════════════════════════════════════════

def main():
    ratings, erreur_agregee, nb_scrutins, scrutins = calculer_ratings()

    # ── Étape 1 : erreurs brutes par scrutin ───────────────────
    print()
    print('═' * 78)
    print('ÉTAPE 1 — ERREURS BRUTES PAR SCRUTIN')
    print('═' * 78)

    for s in scrutins:
        pt          = poids_temporel(s['annee'], ANNEE_REFERENCE, DEMI_VIE_ANS)
        poids_final = s['poids'] * pt
        print()
        print(f"  {s['nom']} ({s['annee']})")
        print(f"  poids scrutin={s['poids']}  ×  poids temporel={pt:.2f}  =  {poids_final:.2f}")
        print(f"  {'Institut':<28} {'Erreur':>8} {'J avant':>8} {'Prox.':>7}")
        print(f"  {'─'*53}")
        for inst, (err, prox) in sorted(s['erreurs'].items(), key=lambda x: x[1][0]):
            jours = round(-DEMI_VIE_JOURS * math.log2(prox)) if prox > 0 else '?'
            print(f"  {inst:<28} {err:>8.2f} {jours:>8} {prox:>7.2f}")

    # ── Étape 2 : rating final par institut ────────────────────
    print()
    print('═' * 78)
    print('ÉTAPE 2 — RATING FINAL PAR INSTITUT')
    print(f'(référence {ANNEE_REFERENCE}, demi-vie {DEMI_VIE_ANS} ans / {DEMI_VIE_JOURS} j, '
          f'minimum {NB_SCRUTINS_MIN} scrutins)')
    print('═' * 78)
    print()
    print(f"  {'Institut':<28} {'Erreur agrégée':>15} {'Rating':>8} {'Scrutins':>9}  ")
    print(f"  {'─'*65}")

    for inst, rating in sorted(ratings.items(), key=lambda x: -x[1]):
        err = erreur_agregee[inst]
        n   = nb_scrutins[inst]
        qualifie = '' if n >= NB_SCRUTINS_MIN else '  ⚠ insuffisant'
        print(f"  {inst:<28} {err:>15.2f} {rating:>8.3f} {n:>9}{qualifie}")

    print()


if __name__ == '__main__':
    main()
