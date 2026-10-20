"""
agregateur.py
─────────────
Calcule la moyenne pondérée des sondages présidentiels 2027 (T1).

Poids de chaque sondage :
  poids = rating_institut × fraîcheur × √min(échantillon, plafond)

La fraîcheur décroît exponentiellement depuis aujourd'hui.
Une série sans plafond est calculée en parallèle, pour comparer les deux
méthodes au résultat réel à l'issue de la campagne.
Le rating est lu depuis calculer_rating.py (même dossier).

Usage :
  python3 agregateur.py
"""

import openpyxl
import math
import csv
import os
from datetime import datetime
from collections import defaultdict

# Source unique du calcul de rating (voir calculer_rating.py)
from calculer_rating import calculer_ratings, normaliser_date

# ══════════════════════════════════════════════════════════════
# PARAMÈTRES — modifier ici si besoin
# ══════════════════════════════════════════════════════════════

FICHIER_SONDAGES = 'sondages_2027.xlsx'
FICHIER_RATING   = 'template_sondages.xlsx'

DATE_REFERENCE   = datetime.today()   # aujourd'hui par défaut
DEMI_VIE_JOURS   = 14                 # fraîcheur : demi-vie en jours

# Plafond de l'échantillon pris en compte dans le poids. Au-delà de quelques
# milliers de répondants, l'erreur d'un sondage tient surtout à la méthode de
# l'institut (redressements, panel), qui ne diminue pas avec la taille : sans
# plafond, √n surpondère les très gros échantillons (ex. Ipsos 09/09/2026,
# n = 13 060). Seuil conventionnel ; entre 2 000 et 5 000, l'effet varie peu.
PLAFOND_ECHANTILLON = 3000
# Les paramètres du rating historique (référence, demi-vies, minimum de
# scrutins) vivent désormais dans calculer_rating.py — source unique.

# Blocs politiques suivis (colonnes du fichier sondages)
BLOCS = ['LFI', 'PS-PP', 'EELV', 'Centre', 'LR', 'RN', 'Reconquête']


# ══════════════════════════════════════════════════════════════
# ÉTAPE 1 — RATINGS (importés de calculer_rating.py)
# ══════════════════════════════════════════════════════════════
# On appelle la fonction de calculer_rating.py avec ses paramètres par
# défaut : les ratings sont donc STRICTEMENT identiques à ceux affichés
# par `python3 calculer_rating.py` (plus de logique dupliquée ni divergente).

ratings, *_ = calculer_ratings(FICHIER_RATING)

# Rating par défaut pour les instituts absents de la base historique
RATING_DEFAUT = min(ratings.values()) if ratings else 0.5


# ══════════════════════════════════════════════════════════════
# ÉTAPE 2 — LECTURE DES SONDAGES 2027
# ══════════════════════════════════════════════════════════════

wb = openpyxl.load_workbook(FICHIER_SONDAGES, data_only=True)
ws = wb['T1 — intentions de vote']

# Repérer les colonnes par leur en-tête
col_index = {}
for c in range(1, ws.max_column + 1):
    val = ws.cell(1, c).value
    if val:
        col_index[val.replace('\n', ' ').strip()] = c

COL_INSTITUT  = col_index.get('Institut')
COL_DATE      = col_index.get('Date terrain (fin)')
COL_ECHAN     = col_index.get('Échantillon')
COL_CANDIDAT  = col_index.get('Candidat testé (Centre)')

sondages = []
for r in range(2, ws.max_row + 1):
    inst = ws.cell(r, COL_INSTITUT).value
    if not inst:
        continue
    date = normaliser_date(ws.cell(r, COL_DATE).value)
    n    = ws.cell(r, COL_ECHAN).value
    if not date or not n:
        continue

    scores = {}
    for bloc in BLOCS:
        c = col_index.get(bloc)
        if c:
            v = ws.cell(r, c).value
            # Un 0 signale un bloc NON testé dans cette hypothèse (ex : PS non
            # mesuré par l'institut), pas un score réel de 0 %. On l'ignore
            # pour ne pas tirer artificiellement la moyenne du bloc vers le bas.
            if isinstance(v, (int, float)) and v > 0:
                scores[bloc] = v

    if not scores:
        continue

    sondages.append({
        'institut': inst,
        'date':     date,
        'n':        n,
        'scores':   scores,
    })

sondages.sort(key=lambda s: s['date'])


# ══════════════════════════════════════════════════════════════
# ÉTAPE 3 — CALCUL DE LA MOYENNE PONDÉRÉE
# ══════════════════════════════════════════════════════════════

def fraicheur(date_sondage, date_ref, demi_vie_jours):
    jours = max((date_ref - date_sondage).days, 0)
    return 0.5 ** (jours / demi_vie_jours)


def calculer_moyennes(sondages, date_ref, plafond=PLAFOND_ECHANTILLON):
    """Moyenne pondérée par bloc des sondages fournis, la fraîcheur étant
    mesurée depuis date_ref. plafond=None calcule la série sans plafond.
    Retourne (moyennes, details) — details sert à l'affichage."""
    numerateur_blocs   = defaultdict(float)
    denominateur_blocs = defaultdict(float)
    details = []

    for s in sondages:
        inst    = s['institut']
        rating  = ratings.get(inst, RATING_DEFAUT)
        fraic   = fraicheur(s['date'], date_ref, DEMI_VIE_JOURS)
        n_retenu = s['n'] if plafond is None else min(s['n'], plafond)
        racine_n = math.sqrt(n_retenu)
        poids   = rating * fraic * racine_n

        for bloc, score in s['scores'].items():
            numerateur_blocs[bloc]   += score * poids
            denominateur_blocs[bloc] += poids

        details.append({**s, 'rating': rating, 'fraicheur': fraic, 'poids': poids})

    moyennes = {
        bloc: numerateur_blocs[bloc] / denominateur_blocs[bloc]
        for bloc in BLOCS
        if denominateur_blocs[bloc] > 0
    }
    return moyennes, details


moyennes, details = calculer_moyennes(sondages, DATE_REFERENCE)
moyennes_sans_plafond, _ = calculer_moyennes(sondages, DATE_REFERENCE, plafond=None)


# ══════════════════════════════════════════════════════════════
# AFFICHAGE
# ══════════════════════════════════════════════════════════════

print()
print('═' * 72)
print('SONDAGES PRIS EN COMPTE')
print('═' * 72)
print()
print(f"  {'Institut':<22} {'Date':>12} {'n':>6} {'Rating':>8} {'Fraîch.':>8} {'Poids':>8}")
print(f"  {'─'*66}")
for d in details:
    print(f"  {d['institut']:<22} {d['date'].strftime('%d/%m/%Y'):>12} "
          f"{d['n']:>6} {d['rating']:>8.3f} {d['fraicheur']:>8.3f} {d['poids']:>8.1f}")

print()
print('═' * 72)
print(f"MOYENNE PONDÉRÉE — T1 PRÉSIDENTIELLE 2027")
print(f"(calculée le {DATE_REFERENCE.strftime('%d/%m/%Y')}, "
      f"demi-vie fraîcheur = {DEMI_VIE_JOURS} jours, "
      f"échantillon plafonné à {PLAFOND_ECHANTILLON})")
print('═' * 72)
print()
for bloc, moy in sorted(moyennes.items(), key=lambda x: -x[1]):
    barre = '█' * int(moy / 2)
    print(f"  {bloc:<14} {moy:>6.1f}%  {barre:<20} "
          f"(sans plafond : {moyennes_sans_plafond[bloc]:.1f}%)")
print()
total = sum(moyennes.values())
print(f"  {'TOTAL':<14} {total:>6.1f}%")
print()


# ══════════════════════════════════════════════════════════════
# ÉTAPE 4 — EXPORT CSV
# ══════════════════════════════════════════════════════════════
# Le fichier est ENTIÈREMENT régénéré à chaque exécution (ce n'est plus un
# journal d'ajouts) : un point par date de fin de terrain distincte, calculé
# avec les seuls sondages connus à cette date. Conséquences :
#   - le script est idempotent : le relancer dix fois donne le même fichier ;
#   - un lot de plusieurs sondages saisis d'un coup produit plusieurs points ;
#   - un sondage saisi après coup corrige son propre point et tous les suivants.
# NE PAS éditer ce fichier à la main : il est écrasé à chaque exécution.
#
# Chaque point est daté à la FIN DE TERRAIN (et non à la date d'exécution) :
# il se place ainsi là où vivait la donnée. Mesurer la fraîcheur depuis cette
# date plutôt que depuis aujourd'hui ne change PAS les pourcentages : la
# fraîcheur ne varie que d'un facteur commun qui se simplifie dans la moyenne.
# (Le nom de colonne 'date_calcul' est conservé pour compatibilité des courbes.)
#
# Deux fichiers, de même format : la série officielle (avec plafond, lue par
# probabilites.py) et la série sans plafond, gardée pour comparer les deux
# méthodes au résultat de l'élection.

FICHIER_CSV              = 'historique_agregateur.csv'
FICHIER_CSV_SANS_PLAFOND = 'historique_agregateur_sans_plafond.csv'
colonnes    = ['date_calcul'] + BLOCS + ['total']

dates_points = sorted({s['date'] for s in sondages})

for fichier, plafond in [(FICHIER_CSV, PLAFOND_ECHANTILLON),
                         (FICHIER_CSV_SANS_PLAFOND, None)]:
    with open(fichier, 'w', newline='', encoding='utf-8') as f:
        writer = csv.DictWriter(f, fieldnames=colonnes)
        writer.writeheader()
        for date_point in dates_points:
            inclus = [s for s in sondages if s['date'] <= date_point]
            moy_point, _ = calculer_moyennes(inclus, date_point, plafond)
            ligne = {'date_calcul': date_point.strftime('%Y-%m-%d')}
            ligne.update({bloc: round(m, 2) for bloc, m in moy_point.items()})
            ligne['total'] = round(sum(moy_point.values()), 2)
            writer.writerow(ligne)

    print(f"  → {fichier} régénéré : {len(dates_points)} points "
          f"(du {dates_points[0]:%d/%m/%Y} au {dates_points[-1]:%d/%m/%Y})")
